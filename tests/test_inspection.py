"""Acceptance checks for inspection evidence, not production PDF generation."""
import hashlib
import json
from pathlib import Path
import unicodedata
import re
import pymupdf as fitz
import openpyxl
import yaml
import pytest

ROOT = Path(__file__).resolve().parents[1]

@pytest.fixture(scope='module')
def evidence():
    return json.loads((ROOT/'tests/data/inspection_evidence.json').read_text(encoding='utf-8'))

def test_readonly_sources(evidence):
    expected=evidence['source_hashes_before']
    actual={p.relative_to(ROOT).as_posix():hashlib.sha256(p.read_bytes()).hexdigest()
            for name in ('input','assets','samples') for p in (ROOT/name).rglob('*') if p.is_file()}
    assert expected == evidence['source_hashes_after'] == actual
    assert evidence['sources_unchanged'] and not evidence['errors']

def test_every_pdf_page_is_accounted_for(evidence):
    expected={p.relative_to(ROOT).as_posix() for folder in ('input','assets','samples') for p in (ROOT/folder).glob('*.pdf')}
    assert set(evidence['pdfs']) == expected
    for name,pages in evidence['pdfs'].items():
        with fitz.open(ROOT/name) as doc:
            assert len(pages)==len(doc)
            for pg,row in zip(doc,pages):
                assert row['page']==pg.number+1
                assert row['rotation']==pg.rotation
                assert row['size']==[round(pg.rect.width,3),round(pg.rect.height,3)]
                assert row['orientation']==('افقی' if pg.rect.width>pg.rect.height else 'عمودی' if pg.rect.width<pg.rect.height else 'مربع')

def test_asset_uniqueness_and_config(evidence):
    cfg=yaml.safe_load((ROOT/'cache/config.yaml').read_text(encoding='utf-8'))
    for role,pattern in cfg['asset_patterns'].items():
        paths=list((ROOT/'assets').glob(pattern))
        assert len(paths)==1
        assert evidence['assets'][role]==[paths[0].relative_to(ROOT).as_posix()]
        assert cfg['resolved_assets'][role]==evidence['assets'][role][0]

def test_workbook_rows_and_headers(evidence):
    wb=openpyxl.load_workbook(ROOT/evidence['assets']['syllabus'][0],read_only=False,data_only=False)
    assert wb.sheetnames==[s['name'] for s in evidence['sheets']]
    for s in evidence['sheets']:
        ws=wb[s['name']]
        assert s['header'] is not None
        assert [c.value for c in ws[s['header']]]==s['headers']
        assert len(s['sample_rows'])==10
        assert len(ws.merged_cells.ranges)==len(s['merged'])
    for m in evidence['matches']:
        ws=wb[m['sheet']]
        assert [c.value for c in ws[m['row']]]==m['values']
        assert ws.cell(m['row'],m['column']).value==m['value']
    assert sorted(m['input'] for m in evidence['matches'])==sorted(p.name for p in (ROOT/'input').glob('*.pdf'))
    wb.close()

def test_required_renders_are_present_and_decodable(evidence):
    expected=sum(len(pages) if name in evidence['assets']['template']+evidence['assets']['border'] else min(2,len(pages)) if name.startswith('samples/') else 1 for name,pages in evidence['pdfs'].items())+1
    assert len(evidence['renders'])==expected
    for item in evidence['renders']:
        pix=fitz.Pixmap(ROOT/item['path'])
        assert pix.width>0 and pix.height>0
        assert hashlib.sha256(pix.samples).hexdigest()==item['sha256_pixels']

def test_font_embedding_and_no_overflow(evidence):
    assert evidence['font_test']['result'][0]>=0
    assert evidence['font_test']['result'][1]==1
    with fitz.open(ROOT/'debug/font_test.pdf') as doc:
        assert len(doc)==1
        assert any(len(doc.extract_font(f[0])[3])>0 for f in doc[0].get_fonts(full=True))

def test_report_has_twelve_sections_and_unknowns(evidence):
    report=(ROOT/'inspect_report.md').read_text(encoding='utf-8')
    assert len(re.findall(r'^## [۰-۹]+\.',report,re.M))==12
    assert 'نامعلوم' in report
    for name in evidence['pdfs']:assert name in report
    for m in evidence['matches']:assert m['value'] in report
    for s in evidence['sheets']:assert s['name'] in report
    visual=json.loads((ROOT/'debug/visual_review.json').read_text(encoding='utf-8'))
    assert visual['source_hashes']==evidence['source_hashes_before']

def test_sample_page_counts_and_orientations(evidence):
    def norm(v):return re.sub(r'[\s_-]+','_',unicodedata.normalize('NFKC',v).casefold())
    for name,ip in evidence['pdfs'].items():
        if not name.startswith('input/'):continue
        candidates=[sp for sn,sp in evidence['pdfs'].items() if sn.startswith('samples/') and norm(Path(name).stem) in norm(Path(sn).stem)]
        assert len(candidates)==1
        sp=candidates[0]
        assert len(sp)==len(ip)+1
        assert [p['orientation'] for p in sp[1:]]==[p['orientation'] for p in ip]
        assert sp[0]['text'].strip()==''
        assert len(sp[0]['images'])==1
