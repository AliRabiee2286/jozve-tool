"""Independent workbook and filename acceptance tests for step 1."""
from pathlib import Path
import copy
import warnings
import tempfile
import openpyxl
import pytest
import yaml
from syllabus import Syllabus, NotFoundError, AmbiguousError, SchemaError, AssetDiscoveryError, normalize_key, load_config, discover_assets
from naamak import resolve, to_output_name

ROOT = Path(__file__).resolve().parents[1]

@pytest.fixture
def tmp_path():
    folder = ROOT / 'tests/data'
    folder.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='step1_', dir=folder) as directory:
        yield Path(directory)

@pytest.fixture(scope='module')
def real():
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        obj=Syllabus(ROOT/'config.yaml')
    assert any('اقتصاد' in str(w.message) for w in caught)
    return obj

@pytest.fixture
def make_book(tmp_path):
    def build(sheets, mapping=None):
        cfg=load_config(ROOT/'config.yaml')
        cfg['sheet_to_template']=mapping or {name:i+1 for i,name in enumerate(sheets)}
        wb=openpyxl.Workbook(); wb.active.title=next(iter(sheets))
        for i,(name,rows) in enumerate(sheets.items()):
            ws=wb.active if i==0 else wb.create_sheet(name)
            for row in rows:ws.append(row)
        p=tmp_path/'case.xlsx';wb.save(p);wb.close()
        return Syllabus.from_workbook(p,cfg)
    return build

H=['نامک','نام استاد','درس','فصل','مبحث']
R=['key_9','استاد','درس الف','فصل الف','مبحث الف']

@pytest.mark.parametrize('namak,sheet,index',[
    ('physics_thermo_8','فیزیک',6),('physics_thermo_9','فیزیک',6),('physics_thermo_10','فیزیک',6),
    ('Savad resane-ma va resane-10','سواد رسانه',7)])
def test_real_exact_independent_cells(real,namak,sheet,index):
    # Oracle reads known workbook coordinates independently of implementation/header detection.
    with open(real.workbook_path,'rb') as f:
        wb=openpyxl.load_workbook(f,data_only=True,read_only=True)
        rows=[row for row in wb[sheet].iter_rows(values_only=True) if len(row)>1 and row[1]==namak]
        assert len(rows)==1
        row=rows[0]
        expected=dict(zip(['namak','teacher','course','chapter','topic'],[v.strip() if isinstance(v,str) else v for v in row[1:6]]))
        expected.update(sheet=sheet,template_index=index)
        assert real.find(namak)==expected
        wb.close()

@pytest.mark.parametrize('name,expected',[
 ('Savad resane-ma va resane-10','Savad resane-ma va resane-10'),
 ('Savad_resane-ma_va_resane-10','Savad resane-ma va resane-10'),
 ('physics_thermo_9 (1)','physics_thermo_9'),('physics_thermo_9__1_','physics_thermo_9'),
 ('physics_thermo_9 (1) (1)','physics_thermo_9'),('physics_thermo_9 2026-10-28','physics_thermo_9'),
 ('physics_thermo_9_20261028','physics_thermo_9'),('2026_10_28_physics_thermo_9.PDF','physics_thermo_9'),
 ('1405-08-07 physics_thermo_9 (1).pdf','physics_thermo_9'),
 ('physics_thermo_9_2026_10_28 (2).pdf','physics_thermo_9'),
 ('physics_thermo_9 (1) 2026-10-28.pdf','physics_thermo_9')])
def test_names(real,name,expected):assert resolve(name,real)==expected

def test_not_found_and_no_numeric_trimming(real):
    for value in ['unknown','physics_thermo_900','physics_thermo_9_2','physics_thermo_9 2026-99-99']:
        with pytest.raises(NotFoundError):resolve(value,real)

@pytest.mark.parametrize('other',[['key_9','دیگری','درس الف','فصل الف','مبحث الف'],['key-9','استاد','درس الف','فصل الف','مبحث الف']])
def test_conflicting_duplicate(make_book,other):
    s=make_book({'رشته':[H,R,other]})
    with pytest.raises(AmbiguousError):s.find('key_9')
    assert s.duplicate_groups

def test_identical_rows_and_ignored_sheet(make_book):
    s=make_book({'رشته':[H,R,R], 'گزارش':[H,['key_9','غلط','غلط','غلط','غلط']]},{'رشته':1})
    assert s.find('key_9')['teacher']=='استاد'
    assert s.duplicate_groups[0]['ambiguous'] is False
    assert s.ignored_sheets==['گزارش']

def test_cross_sheet_duplicate(make_book):
    s=make_book({'الف':[H,R], 'ب':[H,R]})
    with pytest.raises(AmbiguousError):s.find('key_9')

def test_teacher_fallback(make_book):
    with pytest.warns(UserWarning,match='اطلاعات عمومی'):
        s=make_book({'اقتصاد':[['نامک','اطلاعات عمومی','درس','فصل','مبحث'],R]})
    assert s.find('key_9')['teacher']=='استاد'

def test_no_unsafe_fallback(make_book):
    with pytest.raises(SchemaError):make_book({'الف':[['نامک','درس','فصل','مبحث'],['key','x','y','z']]})

def test_reordered_headers_preserve_text(make_book):
    text='  ي ك  رسانه\u200cها  '
    s=make_book({'الف':[['گروه'],[' مبحث ',' نامک ','فصل','استاد','درس'],[text,' کليد-۹ ',text,text,text]]})
    r=s.find('کلید_9')
    assert r['namak']=='کليد-۹'
    assert all(r[k]==text.strip() for k in ('teacher','course','chapter','topic'))
    assert s.column_map['الف']['header_row']==2

@pytest.mark.parametrize('prefix',[4,5])
def test_header_scan_limit(make_book,prefix):
    if prefix==4:assert make_book({'الف':[['مقدمه']]*prefix+[H,R]}).find('key_9')
    else:
        with pytest.raises(SchemaError):make_book({'الف':[['مقدمه']]*prefix+[H,R]})

@pytest.mark.parametrize('headers',[['نامک','استاد','درس','درس','فصل','مبحث'],['نامک','استاد','فصل','مبحث'],['نامک','نامک','استاد','درس','فصل','مبحث']])
def test_bad_schema(make_book,headers):
    with pytest.raises(SchemaError):make_book({'الف':[headers]})

def test_match_first_and_ambiguity_propagates(make_book):
    s=make_book({'الف':[H,R,['key_9 (1)','full','a','b','c'],['key_9 2026-10-28','dated','a','b','c']]})
    assert resolve('key_9 (1).pdf',s)=='key_9 (1)'
    assert resolve('key_9 2026-10-28.pdf',s)=='key_9 2026-10-28'
    s=make_book({'الف':[H,R,['key_9 (1)','a','a','a','a'],['key_9 (1)','b','a','a','a']]})
    with pytest.raises(AmbiguousError):resolve('key_9 (1)',s)

def test_normalization():
    assert normalize_key('  Ａ_ك\u200cي-۱۲٣\u200e ')=='a_کی_123'
    assert normalize_key('ab\ufeff\u200b\u2060cd')=='abcd'

def test_output_name():
    assert to_output_name('Savad resane-ma va resane-10')=='LessonNote_hard_Savad_resane-ma_va_resane-10_v01.pdf'
    assert to_output_name('a<>:"/\\|?*b\x00')=='LessonNote_hard_ab_v01.pdf'
    assert to_output_name('a  b')=='LessonNote_hard_a__b_v01.pdf'
    with pytest.raises(ValueError):to_output_name('<>:')

def test_discovery(tmp_path):
    cfg=load_config(ROOT/'config.yaml');assert len(discover_assets(cfg,ROOT))==4
    cfg=copy.deepcopy(cfg);cfg['paths']['assets']='.';cfg['asset_patterns']={'syllabus':'*Syllabus*.xlsx'}
    with pytest.raises(AssetDiscoveryError):discover_assets(cfg,tmp_path)
    (tmp_path/'Syllabus.xlsx').touch();assert discover_assets(cfg,tmp_path)['syllabus'].name=='Syllabus.xlsx'
    (tmp_path/'Syllabus (1).xlsx').touch()
    with pytest.raises(AssetDiscoveryError):discover_assets(cfg,tmp_path)

def test_explicit_configured_duplicate_headers(make_book):
    # Separate explicit configuration is checked against actual header text.
    import tempfile
    cfg=load_config(ROOT/'config.yaml')
    cfg['sheet_to_template']={'الف':1}
    cfg['header_overrides']={'الف':{'namak':'B','course':'D'}}
    folder=ROOT/'tests/data';folder.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(dir=folder) as directory:
        p=Path(directory)/'overrides.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='الف'
        ws.append(['نامک','نامک','استاد','درس','درس','فصل','مبحث'])
        ws.append([False,'real','t','course',False,'c','topic']);wb.save(p);wb.close()
        s=Syllabus.from_workbook(p,cfg);assert s.find('real')['course']=='course'
        cfg['header_overrides']['الف']['course']='C'
        with pytest.raises(SchemaError):Syllabus.from_workbook(p,cfg)


def test_formula_uses_cached_value_not_formula(make_book):
    s=make_book({'الف':[H,['key','="teacher"','course','chapter','topic']]})
    assert s.find('key')['teacher'] is None  # openpyxl does not calculate formulas


def test_copy_numbers_and_lessons_are_distinct(make_book):
    s=make_book({'الف':[H,['part','t','c','h','p'],['part_9','t','c','h','p'],['part-10','t','c','h','p']]})
    assert resolve('part_9_1_',s)=='part_9'
    assert resolve('part-10 (2)(3).pdf',s)=='part-10'
    with pytest.raises(NotFoundError):resolve('part_11',s)
