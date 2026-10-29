"""Read-only syllabus lookup. No PDF imports, processing, or writes.

Usage: syllabus = Syllabus(); syllabus.find('canonical or normalized namak').
Diagnostics: column_map, warnings, duplicate_groups, ignored_sheets.
"""
from collections import defaultdict
from pathlib import Path
import re
import unicodedata
import warnings
import openpyxl
from openpyxl.utils.cell import column_index_from_string
import yaml

DEFAULT_CONFIG = Path(__file__).with_name('config.yaml')

class NotFoundError(LookupError):
    """No exact normalized syllabus key matched."""

class AmbiguousError(LookupError):
    """A key points to different records; candidates must not be guessed."""
    def __init__(self, namak, matches):
        self.namak, self.matches = namak, matches
        super().__init__(f'Ambiguous namak {namak!r}: {matches!r}')

class SchemaError(ValueError):
    """A configured sheet or its required header is missing/ambiguous."""

class AssetDiscoveryError(ValueError):
    """An asset pattern did not match exactly one file."""


def normalize_key(value):
    """Normalization is for lookup only, never for display values."""
    if not isinstance(value, str):
        raise TypeError('Lookup keys must be strings')
    text = unicodedata.normalize('NFKC', value).translate(str.maketrans('يك', 'یک'))
    text = ''.join(c for c in text if unicodedata.category(c) != 'Cf'
                   and not (unicodedata.category(c) == 'Cc' and not c.isspace())
                   and ord(c) != 0x034F and not 0xFE00 <= ord(c) <= 0xFE0F
                   and not 0xE0100 <= ord(c) <= 0xE01EF)
    text = ''.join(str(unicodedata.decimal(c)) if c.isdecimal() else c for c in text)
    return re.sub(r'[\s_-]+', '_', text.lower().strip()).strip('_')


def load_config(path=DEFAULT_CONFIG):
    with Path(path).open(encoding='utf-8-sig') as stream:
        config = yaml.safe_load(stream)
    if not isinstance(config, dict):
        raise SchemaError('Configuration must be a mapping')
    return config


def discover_assets(config, base_dir):
    """Resolve patterns relative to config, never trust stale resolved filenames."""
    directory = Path(base_dir) / config['paths']['assets']
    result = {}
    for role, pattern in config['asset_patterns'].items():
        found = sorted(p for p in directory.glob(pattern) if p.is_file())
        if len(found) != 1:
            raise AssetDiscoveryError(f'{role}: pattern {pattern!r}: {len(found)} files')
        result[role] = found[0].resolve()
    return result


def _trim(value):
    return value.strip() if isinstance(value, str) else value


class Syllabus:
    def __init__(self, config_path=DEFAULT_CONFIG):
        path = Path(config_path).resolve()
        config = load_config(path)
        self.assets = discover_assets(config, path.parent)
        self._read(self.assets['syllabus'], config)

    @classmethod
    def from_workbook(cls, workbook_path, config):
        """Explicit workbook entry point for isolated tests and offline callers."""
        obj = cls.__new__(cls)
        obj.assets = {}
        obj._read(Path(workbook_path), config)
        return obj

    def _read(self, path, config):
        self.workbook_path = Path(path)
        self.column_map, self.warnings = {}, []
        self._index = defaultdict(list)
        mapping = config['sheet_to_template']
        if not mapping or any(type(v) is not int or v < 1 for v in mapping.values()):
            raise SchemaError('Template indices must be positive one-based integers')
        aliases = {key: {normalize_key(v) for v in names} for key, names in config['headers'].items()}
        if set(aliases) != {'namak', 'teacher', 'course', 'chapter', 'topic'}:
            raise SchemaError('Configuration must define all five semantic headers')
        scan = config['header_scan_rows']
        if type(scan) is not int or not 1 <= scan <= 5:
            raise SchemaError('header_scan_rows must be between 1 and 5')
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            missing = set(mapping) - set(wb.sheetnames)
            if missing:
                raise SchemaError(f'Missing configured sheets: {sorted(missing)}')
            self.ignored_sheets = [s for s in wb.sheetnames if s not in mapping]
            for name, template in mapping.items():
                ws = wb[name]
                top = list(ws.iter_rows(min_row=1, max_row=scan, values_only=True))
                overrides = config.get('header_overrides', {}).get(name, {})
                if set(overrides) - set(aliases):
                    raise SchemaError(f'{name}: unknown header override')
                hints = {field: column_index_from_string(col) for field, col in overrides.items()}
                positions = [(ri, ci) for ri, row in enumerate(top, 1) for ci, v in enumerate(row, 1)
                             if isinstance(v, str) and normalize_key(v) in aliases['namak']
                             and ('namak' not in hints or ci == hints['namak'])]
                if len(positions) != 1:
                    raise SchemaError(f'{name}: expected exactly one namak header in first {scan} rows; {positions}')
                header_row, namak_col = positions[0]
                header = top[header_row - 1]
                columns = {'namak': namak_col}
                for field in ('course', 'chapter', 'topic', 'teacher'):
                    found = [ci for ci, v in enumerate(header, 1)
                             if isinstance(v, str) and normalize_key(v) in aliases[field]
                             and (field not in hints or ci == hints[field])]
                    if len(found) > 1 or (not found and (field != 'teacher' or field in hints)):
                        raise SchemaError(f'{name}: missing or duplicate {field} header: {found}')
                    if found:
                        columns[field] = found[0]
                fallback = 'teacher' not in columns
                if fallback:
                    candidate = namak_col + 1
                    forbidden = set().union(*(aliases[k] for k in ('course', 'chapter', 'topic', 'namak')))
                    if candidate > len(header) or (isinstance(header[candidate - 1], str)
                                                  and normalize_key(header[candidate - 1]) in forbidden):
                        raise SchemaError(f'{name}: unsafe teacher fallback after namak')
                    columns['teacher'] = candidate
                    message = f'{name}: teacher fallback to column {candidate}, header {header[candidate - 1]!r}'
                    self.warnings.append(message)
                    warnings.warn(message, UserWarning, stacklevel=3)
                if len(set(columns.values())) != len(columns):
                    raise SchemaError(f'{name}: multiple fields share a column')
                self.column_map[name] = {'header_row': header_row, 'columns': columns,
                                         'headers': {k: header[v - 1] for k, v in columns.items()},
                                         'teacher_fallback': fallback}
                for ri, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                    namak = _trim(row[namak_col - 1])
                    if namak is None or namak == '':
                        continue
                    if not isinstance(namak, str):
                        raise SchemaError(f'{name}!{ri}: namak must be text')
                    key = normalize_key(namak)
                    if not key:
                        continue
                    record = {k: _trim(row[col - 1]) for k, col in columns.items()}
                    record.update(sheet=name, template_index=template)
                    self._index[key].append({'row': ri, 'record': record})
        finally:
            wb.close()
        self.duplicate_groups = [dict(key=k, ambiguous=self._ambiguous(rows), matches=rows)
                                 for k, rows in sorted(self._index.items()) if len(rows) > 1]

    @staticmethod
    def _ambiguous(rows):
        return any(item['record'] != rows[0]['record'] for item in rows[1:])

    def find(self, namak):
        rows = self._index.get(normalize_key(namak), [])
        if not rows:
            raise NotFoundError(f'No syllabus match: {namak!r}')
        if self._ambiguous(rows):
            raise AmbiguousError(namak, rows)
        return dict(rows[0]['record'])
